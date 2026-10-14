# -*- coding: utf-8 -*-
"""Prova OFFLINE dos VALORES FIXOS do sistema (Empresa/Coordenador/Horário).

Nada de Playwright, nada de formulário ao vivo: testa só a lógica Python de
declarar os fixos e injetá-los em TODA combinação da planilha. Sai 0 se tudo ok.
"""
import io
import os
import sys
import json
import tempfile

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from core.forms_submitter import FormsSubmitter

falhas = []


def ok(cond, msg):
    print(("  OK  " if cond else " FALHA ") + msg)
    if not cond:
        falhas.append(msg)


def _planilha(path, header, linhas):
    wb = Workbook(); ws = wb.active
    ws.append(header)
    for ln in linhas:
        ws.append(ln)
    wb.save(path)


tmp = tempfile.mkdtemp(prefix="fixos_")

# ---------------------------------------------------------------- #
print("\n[1] valores_fixos() padrão = constantes declaradas")
fx = {f["pergunta"]: f["valor"] for f in FormsSubmitter.valores_fixos(config_path=os.path.join(tmp, "nao_existe.json"))}
ok(fx.get("Empresa") == "Guaraná Diesel", f"Empresa = {fx.get('Empresa')!r}")
ok(fx.get("Coordenador") == "Bruno Duarte", f"Coordenador = {fx.get('Coordenador')!r}")
ok(fx.get("Horário") == "diurno", f"Horário = {fx.get('Horário')!r}")

# ---------------------------------------------------------------- #
print("\n[2] planilha LARGA multi-dia SEM as 3 colunas -> fixos em TODA combinação")
largo = os.path.join(tmp, "largo.xlsx")
_planilha(
    largo,
    ["URL", "https://forms.exemplo/x", "", ""],  # linha de URL
    [],
)
# refaz com cabeçalho de perguntas + 3 dias (só o que VARIA)
wb = Workbook(); ws = wb.active
ws.append(["URL", "https://forms.exemplo/x"])
ws.append(["Km inicial", "Km final", "Tipo de interdição", "Serviço"])
ws.append([939, 42, "N/A", "Limpeza em Geral"])
ws.append([42, 86, "N/A", "Limpeza em Geral"])
ws.append([120, 130, "Estreitamento", "Conserva Áreas Verdes"])
wb.save(largo)

plano = FormsSubmitter.parse_reference_sheet_multi(largo)
combos = FormsSubmitter.aplicar_valores_fixos(plano["combinacoes"],
                                               fixos=FormsSubmitter.valores_fixos(
                                                   config_path=os.path.join(tmp, "nao_existe.json")))
ok(plano["url"] == "https://forms.exemplo/x", f"url lida = {plano['url']!r}")
ok(len(combos) == 3, f"3 combinações (dias) -> {len(combos)}")
for i, c in enumerate(combos, 1):
    d = {x["pergunta"]: x["valor"] for x in c}
    tem = (d.get("Empresa") == "Guaraná Diesel" and d.get("Coordenador") == "Bruno Duarte"
           and d.get("Horário") == "diurno")
    ok(tem, f"combinação {i} tem os 3 fixos (e manteve o que varia: {d.get('Serviço')!r})")

print("\n   amostra da combinação 1 (o que a Sophia vai preencher):")
for x in combos[0]:
    print(f"     - {x['pergunta']}: {x['valor']}")

# ---------------------------------------------------------------- #
print("\n[3] planilha COM coluna 'Empresa' errada -> fixo SOBRESCREVE, sem duplicar")
wb = Workbook(); ws = wb.active
ws.append(["Empresa", "Km inicial", "Km final", "Serviço"])
ws.append(["Empresa Errada SA", 10, 20, "Limpeza em Geral"])
sobre = os.path.join(tmp, "sobre.xlsx")
wb.save(sobre)
plano = FormsSubmitter.parse_reference_sheet_multi(sobre)
combos = FormsSubmitter.aplicar_valores_fixos(plano["combinacoes"],
                                               fixos=FormsSubmitter.valores_fixos(
                                                   config_path=os.path.join(tmp, "nao_existe.json")))
c = combos[0]
empresas = [x["valor"] for x in c if x["pergunta"].strip().lower() == "empresa"]
ok(empresas == ["Guaraná Diesel"], f"Empresa única e fixa -> {empresas}")
ok(all(x["valor"] != "Empresa Errada SA" for x in c), "valor errado da planilha foi descartado")

# ---------------------------------------------------------------- #
print("\n[4] override por config.json (sem recompilar) — dict e lista")
cfg_dict = os.path.join(tmp, "config_dict.json")
with open(cfg_dict, "w", encoding="utf-8") as f:
    json.dump({"forms_valores_fixos": {"Coordenador": "Outro Chefe", "Horário": "noturno",
                                        "Empresa": "Guaraná Diesel"}}, f, ensure_ascii=False)
fx = {f["pergunta"]: f["valor"] for f in FormsSubmitter.valores_fixos(config_path=cfg_dict)}
ok(fx.get("Coordenador") == "Outro Chefe" and fx.get("Horário") == "noturno",
   f"config dict respeitado -> Coordenador={fx.get('Coordenador')!r}, Horário={fx.get('Horário')!r}")

cfg_list = os.path.join(tmp, "config_list.json")
with open(cfg_list, "w", encoding="utf-8") as f:
    json.dump({"forms_valores_fixos": [{"pergunta": "Empresa", "valor": "Guaraná Diesel"},
                                       {"pergunta": "Coordenador", "valor": "Bruno Duarte"},
                                       {"pergunta": "Horário", "valor": "diurno"}]}, f, ensure_ascii=False)
fx = {f["pergunta"]: f["valor"] for f in FormsSubmitter.valores_fixos(config_path=cfg_list)}
ok(fx.get("Empresa") == "Guaraná Diesel" and fx.get("Horário") == "diurno",
   "config lista respeitado")

# ---------------------------------------------------------------- #
print("\n[5] config inválido/corrompido -> cai nos padrões (nunca quebra)")
cfg_bad = os.path.join(tmp, "config_bad.json")
with open(cfg_bad, "w", encoding="utf-8") as f:
    f.write("{ isto nao e json valido ")
fx = {f["pergunta"]: f["valor"] for f in FormsSubmitter.valores_fixos(config_path=cfg_bad)}
ok(fx.get("Empresa") == "Guaraná Diesel", "config quebrado -> padrões")

print("\n" + ("TODOS OS TESTES PASSARAM ✅" if not falhas else f"{len(falhas)} FALHA(S) ❌"))
sys.exit(0 if not falhas else 1)
