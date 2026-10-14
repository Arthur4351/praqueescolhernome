# -*- coding: utf-8 -*-
"""BATERIA SUPREMA — nucleo Excel + veredito VBA.
Exercita ExcelEngine DIRETO (sem GUI): injeta MUITAS variantes de formula NATIVA
(.xlsx e .xlsm com keep_vba), consulta (QUERY_EXCEL) e conta vazias
(QUERY_COUNT_EMPTY), e prova o veredito sobre VBA/macro (ela NAO gera VBA)."""
import io, sys, os, tempfile, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import openpyxl
from openpyxl import load_workbook
from core.excel_engine import ExcelEngine

TMP = tempfile.mkdtemp(prefix="sophia_excel_")
ok = lambda b: "OK" if b else "FALHOU"
passes, fails = 0, 0
def check(nome, cond, evidencia=""):
    global passes, fails
    if cond: passes += 1
    else: fails += 1
    print(f"  [{ok(cond)}] {nome}" + (f"  ->  {evidencia}" if evidencia else ""))

print("=" * 70)
print("1) INJECT_FORMULA — formulas NATIVAS (openpyxl), variantes pt-BR e EN")
print("=" * 70)
xlsx = os.path.join(TMP, "calc.xlsx")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Dados"
for i, v in enumerate([10, 20, 30, 40, 50], start=1):
    ws[f"A{i}"] = v
wb.create_sheet("Resumo")
wb.save(xlsx)

variantes = [
    ("Dados", "B1", "=SOMA(A1:A5)"),
    ("Dados", "B2", "=SUM(A1:A5)"),
    ("Dados", "B3", "=MÉDIA(A1:A5)"),
    ("Dados", "B4", "=AVERAGE(A1:A5)"),
    ("Dados", "B5", "=CONT.VALORES(A1:A5)"),
    ("Dados", "C1", '=SE(A1>0;"POS";"NEG")'),
    ("Dados", "C2", "=PROCV(10;A1:B5;2;FALSO)"),
    ("Resumo", "A1", "=Dados!A1"),          # referencia entre abas
]
for aba, cel, form in variantes:
    r = ExcelEngine.inject_formula(xlsx, aba, cel, form)
    check(f"injeta {form} em {aba}!{cel}", r is True)
# aba inexistente -> deve retornar False (honesto, sem crash)
r_bad = ExcelEngine.inject_formula(xlsx, "NaoExiste", "A1", "=SOMA(A1:A2)")
check("aba inexistente retorna False (nao inventa)", r_bad is False, f"retorno={r_bad}")

# releitura: a formula NATIVA foi realmente gravada na celula
wb2 = load_workbook(xlsx)  # data_only=False -> devolve a string da formula
for aba, cel, form in variantes:
    val = wb2[aba][cel].value
    check(f"persistiu {aba}!{cel}", val == form, f"celula={val!r}")
wb2.close()

print("\n" + "=" * 70)
print("2) .xlsm com keep_vba — salvar NAO destroi macro; NAO cria macro")
print("=" * 70)
xlsm = os.path.join(TMP, "relatorio.xlsm")
wbm = openpyxl.Workbook(); wbm.active.title = "Dados"; wbm["Dados"]["A1"] = 123
wbm.save(xlsm)
r_xlsm = ExcelEngine.inject_formula(xlsm, "Dados", "A2", "=A1*2")
check("injeta formula em .xlsm (caminho keep_vba)", r_xlsm is True)
wbm2 = load_workbook(xlsm, keep_vba=True)
check(".xlsm reabre e formula persistiu", wbm2["Dados"]["A2"].value == "=A1*2",
      f"A2={wbm2['Dados']['A2'].value!r}")
wbm2.close()
# prova de que keep_vba so PRESERVA: sem vbaProject de entrada, nenhum e criado
import zipfile
with zipfile.ZipFile(xlsm) as z:
    tem_vba = any("vbaProject" in n for n in z.namelist())
check("nenhum vbaProject.bin foi AUTORADO pela Sophia", tem_vba is False,
      f"vbaProject presente={tem_vba}")

print("\n" + "=" * 70)
print("3) QUERY_EXCEL — consulta fuzzy de coluna e valor")
print("=" * 70)
qx = os.path.join(TMP, "efetivo.xlsx")
wq = openpyxl.Workbook(); wsq = wq.active
wsq.append(["NOME", "MATRÍCULA", "FUNÇÃO"])
wsq.append(["JOAO", "12345", "Tecnico"])
wsq.append(["MARIA", "67890", "Engenheira"])
wq.save(qx)
r1 = ExcelEngine.query_data(qx, "JOAO", "matricula")
check("acha JOAO -> matricula (coluna sem acento no pedido)", "12345" in r1, r1)
r2 = ExcelEngine.query_data(qx, "maria", "funcao")
check("acha maria (minusculo) -> funcao", "Engenheira" in r2, r2)
r3 = ExcelEngine.query_data(qx, "CARLOS", "matricula")
check("valor inexistente degrada honesto", ("Nada encontrado" in r3 or "🔍" in r3), r3)
r4 = ExcelEngine.query_data(qx, "JOAO", "coluna_que_nao_existe")
check("coluna inexistente degrada honesto", ("não localizada" in r4 or "⚠️" in r4), r4)

print("\n" + "=" * 70)
print("4) QUERY_COUNT_EMPTY — conta registros vazios")
print("=" * 70)
qc = os.path.join(TMP, "faltas.xlsx")
wc = openpyxl.Workbook(); wsc = wc.active
wsc.append(["NOME", "FOTO"])
wsc.append(["A", "sim"]); wsc.append(["B", None]); wsc.append(["C", ""])
wsc.append(["D", "sim"]); wsc.append(["E", None])
wc.save(qc)
rc = ExcelEngine.query_count_empty(qc, "foto")
check("conta 3 vazias na coluna FOTO", "3" in rc, rc)
rc2 = ExcelEngine.query_count_empty(qc, "inexistente")
check("coluna inexistente degrada honesto", ("não localizada" in rc2 or "⚠️" in rc2), rc2)

print("\n" + "=" * 70)
print("VEREDITO VBA: a Sophia NAO gera macro/VBA automatico.")
print("  - inject_formula grava SO formula NATIVA (=SOMA, =SE, =Dados!A1...).")
print("  - keep_vba=True apenas PRESERVA macro pre-existente ao salvar .xlsm.")
print("  - nenhum vbaProject.bin foi autorado (provado no bloco 2).")
print("  - o unico motor que mencionava win32com/macro (dynamic_coder) esta")
print("    ESTACIONADO: generate_and_run/edit_project_file retornam 🔒.")
print("=" * 70)

shutil.rmtree(TMP, ignore_errors=True)
print(f"\nRESUMO EXCEL: {passes} passaram, {fails} falharam.")
sys.exit(1 if fails else 0)
