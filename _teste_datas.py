# -*- coding: utf-8 -*-
import sys, os
from pathlib import Path
try: sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception: pass
ROOT = Path(r"C:\Users\arthur 2\Desktop\praqueescolhernome-main"); os.chdir(str(ROOT)); sys.path.insert(0, str(ROOT))
from openpyxl import Workbook, load_workbook
from core.agent_core import SophiaAgentCore
from minhas_habilidades import SophiaExecutor

BASE = ROOT / "_teste_interface_out"; plan = BASE / "datas_populada.xlsx"
wb = Workbook(); ws = wb.active; ws.title = "EQUIPE 01"
ws["A1"]="DATA"; ws["B1"]="ATIVIDADE"
for r,(atv) in enumerate(["Ronda matinal","Vistoria portao","Troca de turno","Ronda noturna"], start=2):
    ws.cell(row=r, column=2).value = atv       # linhas COM conteudo, coluna DATA vazia
wb.save(plan)
print("ANTES  -> DATA A2..A5 =", [ws['A%d'%r].value for r in range(2,6)])

ag = SophiaAgentCore(); ex = SophiaExecutor(autonomous_mode=True)
def log_cb(m): print("        │", m)
print("\n👤 VOCE  : sincroniza as datas dessa planilha de atividades")
resp = ag.evaluate_chat("sincroniza as datas dessa planilha de atividades", "Arthur")
print("🤖 SOPHIA:", resp)
d = ag.pending_intents[0] if ag.pending_intents else {}
print("        └─ NLU:", {k:d.get(k) for k in ('intent','status','score')})
ex.atualizar_datas_planilha(str(plan), str(BASE/"fotos"), log_cb, None)
wb2 = load_workbook(plan); ws2 = wb2["EQUIPE 01"]
print("        📅 DEPOIS -> DATA A2..A5 =", [ws2['A%d'%r].value for r in range(2,6)])
