# -*- coding: utf-8 -*-
"""Teste funcional REAL (headless) do motor da Sophia:
   1) cria planilha de relatorio fotografico, gera fotos e injeta (processar_comando)
   2) cria efetivo (pastas + planilha) e roda a auditoria (auditar_efetivo)
   Nao usa LLM, nao gera codigo, nao acessa rede com credenciais."""
import sys, os, traceback
from pathlib import Path

# Console do Windows e cp1252; a engine loga com emojis -> forca UTF-8 p/ nao quebrar o log
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = Path(r"C:\Users\arthur 2\Desktop\praqueescolhernome-main")
os.chdir(str(ROOT))
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

BASE = ROOT / "_teste_funcional"
FOTOS = BASE / "fotos"
SAIDA = BASE / "saida_relatorios"      # a "pasta especifica" onde o relatorio e salvo
EFETIVO = BASE / "efetivo"
for d in (BASE, FOTOS, SAIDA, EFETIVO):
    d.mkdir(parents=True, exist_ok=True)

log = []
_logfile = BASE / "resultado_teste.txt"
def L(msg):
    line = str(msg)
    log.append(line)
    try:
        print(line, flush=True)
    except Exception:
        print(line.encode("ascii", "replace").decode("ascii"), flush=True)
    with open(_logfile, "a", encoding="utf-8") as fh:
        fh.write(line + "\n")

open(_logfile, "w", encoding="utf-8").close()  # zera o log a cada execucao

EQUIPES = ["EQUIPE 01", "EQUIPE 02"]

# 1) Planilha de relatorio fotografico: uma aba por equipe, cabecalho DATA + colunas de foto
plan_fotos = BASE / "relatorio_fotografico.xlsx"
wb = Workbook(); wb.remove(wb.active)
for eq in EQUIPES:
    ws = wb.create_sheet(eq[:31])
    ws["A1"] = "DATA"; ws["B1"] = 1
    ws["C1"] = "ENTRADA"; ws["D1"] = "SAIDA"; ws["E1"] = "ANTES"; ws["F1"] = "DEPOIS"
    ws.row_dimensions[2].height = 130
    for col in "CDEF":
        ws.column_dimensions[col].width = 22
wb.save(plan_fotos)
L(f"[SETUP] Planilha criada: {plan_fotos.name} (abas: {', '.join(EQUIPES)})")

# 2) Fotos de teste (dia 01, tipos E/S/A/D) por equipe
def make_photo(path, texto):
    img = Image.new("RGB", (640, 480), (70, 110, 160))
    d = ImageDraw.Draw(img)
    d.rectangle([10, 10, 630, 470], outline=(255, 255, 255), width=4)
    d.text((40, 220), texto, fill=(255, 255, 255))
    img.save(path, "JPEG", quality=80)

for eq in EQUIPES:
    pasta = FOTOS / eq
    pasta.mkdir(parents=True, exist_ok=True)
    for suf, nome in [("E", "ENTRADA"), ("S", "SAIDA"), ("A", "ANTES"), ("D", "DEPOIS")]:
        make_photo(pasta / f"01 {suf}.jpg", f"{eq} DIA01 {nome}")
L(f"[SETUP] Fotos geradas: {sum(1 for _ in FOTOS.rglob('*.jpg'))} arquivos em {len(EQUIPES)} pastas")

# 3) Motor real de injecao de fotos
try:
    from minhas_habilidades import SophiaExecutor
    ex = SophiaExecutor(autonomous_mode=False)
    L(f"[INFO] SophiaExecutor instanciado. OCR disponivel? {ex._ocr_disponivel}")
except Exception:
    L("[EXCECAO ao instanciar SophiaExecutor]\n" + traceback.format_exc())
    ex = None

if ex:
    L("\n===== TESTE 1: SALVAR FOTOS NA PLANILHA (processar_comando) =====")
    try:
        ret = ex.processar_comando(str(FOTOS), str(plan_fotos), str(SAIDA), "TESTE", L)
        L(f"[RETORNO] {ret}")
    except Exception:
        L("[EXCECAO]\n" + traceback.format_exc())

    saidas = list(SAIDA.glob("Relatorio_*.xlsm"))
    L(f"[VERIFICA] Arquivos gerados na pasta de saida: {[p.name for p in saidas]}")
    for p in saidas:
        try:
            wbx = load_workbook(p)
            tot = sum(len(getattr(wbx[s], '_images', [])) for s in wbx.sheetnames)
            L(f"[VERIFICA] {p.name}: {tot} imagem(ns) embutida(s) em {len(wbx.sheetnames)} aba(s)")
            wbx.close()
        except Exception as e:
            L(f"[VERIFICA] Falha ao reabrir {p.name}: {e}")

    # 4) Efetivo real
    L("\n===== TESTE 2: EFETIVO / AUDITORIA DE PRESENCA (auditar_efetivo) =====")
    for nome in ["JOAO SILVA", "MARIA SOUZA", "CARLOS PEREIRA"]:
        (EFETIVO / nome).mkdir(parents=True, exist_ok=True)
    plan_efetivo = BASE / "efetivo_lista.xlsx"
    wbe = Workbook(); wse = wbe.active
    wse["A1"] = "NOME"
    for i, nome in enumerate(["JOAO SILVA", "MARIA SOUZA", "ANA LIMA"], start=2):
        wse.cell(row=i, column=1).value = nome
    wbe.save(plan_efetivo)
    L("[SETUP] Pastas efetivo: JOAO/MARIA/CARLOS | Excel: JOAO/MARIA/ANA (ANA so no Excel, CARLOS so na pasta)")
    try:
        ex.auditar_efetivo(str(EFETIVO), str(plan_efetivo), L)
    except Exception:
        L("[EXCECAO]\n" + traceback.format_exc())

L("\n===== FIM DO TESTE FUNCIONAL =====")
