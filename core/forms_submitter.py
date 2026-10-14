"""
core/forms_submitter.py

Preenchedor de formulários (Microsoft Forms / formulários web) guiado por uma
PLANILHA DE REFERÊNCIA. A planilha diz, linha a linha, QUAL pergunta preencher e
O QUE escrever ou marcar; ao final é gerado um PDF do formulário preenchido.

AUTOSSUFICIENTE NO PC CORPORATIVO:
- Usa o Microsoft Edge JÁ instalado no Windows (channel="msedge"): NÃO baixa
  Chromium nem exige instalar navegador. O driver do Playwright é embutido no .exe
  em tempo de build (como o Tesseract), não em tempo de execução.
- Se Playwright/Edge não estiverem disponíveis, degrada com mensagem honesta (sem crash).
- O import do Playwright é adiado para dentro dos métodos: importar este módulo
  NUNCA quebra a aplicação, mesmo sem o Playwright instalado.

Formato da planilha de referência (1ª aba):
  Coluna A: PERGUNTA  (título da pergunta como aparece no formulário)
  Coluna B: TIPO      (opcional: texto|data|opcao|dropdown|checkbox — autodetecta se vazio)
  Coluna C: VALOR     (o que escrever, ou o rótulo da opção a marcar; várias com ';')
Linha especial: PERGUNTA em {URL, LINK, FORMULARIO} -> VALOR = link do formulário.
"""
import os
import re
import json
import unicodedata
from datetime import datetime

# Chaves que, na coluna PERGUNTA, indicam a URL do formulário em vez de um campo.
_URL_KEYS = {"URL", "LINK", "FORMULARIO", "FORMULÁRIO", "FORM"}
# Caminhos padrão do Microsoft Edge (componente de 1ª parte do Windows corporativo).
_EDGE_PATHS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]

# ---------------------------------------------------------------------------
# VALORES FIXOS DO SISTEMA (constantes declaradas pelo usuário).
# São campos que NUNCA mudam entre dias/combinações e por isso NÃO precisam
# aparecer na planilha de referência (que carrega só o que varia: Km, tipo de
# interdição, serviço, data...). Sophia aplica estes valores em TODA combinação.
#
# São SOBRESCRevíveis sem recompilar: basta uma chave "forms_valores_fixos" no
# config.json (ao lado do .exe). Pode ser:
#   • lista: [{"pergunta": "...", "valor": "...", "tipo": ""}, ...]
#   • dict : {"Empresa": "Guaraná Diesel", "Coordenador": "Bruno Duarte", ...}
# Sem a chave, valem os padrões abaixo.
_VALORES_FIXOS_PADRAO = [
    {"pergunta": "Empresa",     "tipo": "", "valor": "Guaraná Diesel"},
    {"pergunta": "Coordenador", "tipo": "", "valor": "Bruno Duarte"},
    {"pergunta": "Horário",     "tipo": "", "valor": "diurno"},
    {"pergunta": "Revisão de Solicitação", "tipo": "", "valor": "Não"},
    {"pergunta": "Contorno",    "tipo": "", "valor": "N/A"},
]


def _norm_pergunta(s):
    """Normaliza um rótulo de pergunta p/ comparação: sem acento, minúsculo, sem ':'."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return s.lower().strip().rstrip(":").strip()


def _colide(pergunta, alvos_norm):
    """True se o rótulo colide (por substring normalizada) com algum alvo."""
    n = _norm_pergunta(pergunta)
    if not n:
        return False
    return any(n == a or n in a or a in n for a in alvos_norm if a)


class FormsSubmitter:
    def __init__(self, output_dir="relatorios"):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir, exist_ok=True)

    # ------------------------------------------------------------------ #
    # Disponibilidade (degrada com honestidade, nunca quebra o app)
    # ------------------------------------------------------------------ #
    @staticmethod
    def is_available():
        """Retorna (ok: bool, motivo: str). Não importa o Playwright de fato."""
        import importlib.util
        if importlib.util.find_spec("playwright") is None:
            return False, "Playwright não disponível — recurso de Forms desativado nesta instalação."
        if not any(os.path.isfile(p) for p in _EDGE_PATHS):
            return False, "Microsoft Edge não encontrado neste computador."
        return True, "OK"

    # ------------------------------------------------------------------ #
    # Leitura da planilha de referência (Python puro — testável offline)
    # ------------------------------------------------------------------ #
    @staticmethod
    def parse_reference_sheet(xlsx_path):
        """Lê a planilha de referência e devolve {'url': str|None, 'campos': [{pergunta,tipo,valor}]}.

        Autodetecta DOIS layouts verticais (sem exigir formato fixo nem nomes de campo):
          • 2 colunas:  CAMPO | VALOR             (o tipo é autodetectado no preenchimento)
          • 3 colunas:  PERGUNTA | TIPO | VALOR   (tipo explícito, se reconhecido)
        Linha especial CAMPO ∈ {URL, LINK, FORMULARIO, FORM} define a URL do formulário.
        """
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        ws = wb.active
        # materializa só as linhas com um "campo" na 1ª coluna (iter_rows é gerador de passada única)
        linhas = [r for r in ws.iter_rows(values_only=True)
                  if r and r[0] is not None and str(r[0]).strip()]
        wb.close()

        # Tokens que valem como TIPO explícito (senão a coluna do meio é tratada como VALOR).
        _TIPOS = {"texto", "text", "data", "date", "numero", "número", "number",
                  "opcao", "opção", "radio", "escolha", "unica", "única",
                  "checkbox", "caixa", "multipla", "múltipla", "marcar",
                  "dropdown", "lista", "select", "combo", "combobox"}

        # Layout de 3 colunas só quando a coluna C tem conteúdo em ALGUMA linha de dados.
        tem_col_c = any(len(r) > 2 and r[2] is not None and str(r[2]).strip() for r in linhas)

        url = None
        campos = []
        for row in linhas:
            pergunta = str(row[0]).strip()
            b = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            c = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            chave = pergunta.upper().rstrip(":").strip()

            # linha de configuração da URL (o link pode estar em B ou C)
            if chave in _URL_KEYS:
                url = c or b or None
                url = str(url).strip() if url else None
                continue
            # ignora um cabeçalho literal "PERGUNTA | TIPO | VALOR"
            if chave == "PERGUNTA" and b.upper() in ("TIPO", "VALOR", "RESPOSTA", ""):
                continue

            if tem_col_c:
                # 3 colunas: B = tipo (só se reconhecido), C = valor
                tipo = b.lower() if b.lower() in _TIPOS else ""
                valor = c
            else:
                # 2 colunas: B = valor; tipo fica vazio e é autodetectado no preenchimento
                tipo = ""
                valor = b
            campos.append({"pergunta": pergunta, "tipo": tipo, "valor": valor})
        return {"url": url, "campos": campos}

    @staticmethod
    def parse_reference_sheet_multi(xlsx_path):
        """Lê a planilha e devolve VÁRIAS combinações (1 preenchimento do form por combinação).

        Autodetecta, além dos layouts verticais (1 combinação), o layout LARGO/tabular:
          • LARGO: 1ª linha = TÍTULOS das perguntas (uma por coluna); cada linha
            seguinte = 1 dia/combinação. É o formato do uso real (vários dias,
            ~N combinações). Detectado quando ALGUMA linha tem ≥4 colunas com dado
            (os verticais têm no máx. 3 colunas: CAMPO|TIPO|VALOR).
          • VERTICAL: reaproveita parse_reference_sheet -> UMA combinação.
        Retorna {'url': str|None, 'combinacoes': [[{pergunta,tipo,valor}, ...], ...]}.
        Nada é hardcodado: perguntas e valores vêm da própria planilha.
        """
        from openpyxl import load_workbook

        def _txt(x):
            return str(x).strip() if x is not None and str(x).strip() else ""

        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        ws = wb.active
        grade = [list(r) for r in ws.iter_rows(values_only=True)
                 if r and any(c is not None and str(c).strip() for c in r)]
        wb.close()

        # separa uma eventual linha de URL (não é pergunta nem dado)
        url = None
        linhas = []
        for r in grade:
            chave = _txt(r[0]).upper().rstrip(":").strip()
            if chave in _URL_KEYS:
                for c in r[1:]:
                    if _txt(c):
                        url = _txt(c)
                        break
                continue
            linhas.append(r)

        # LARGO se alguma linha usa uma 4ª coluna (verticais nunca passam de 3)
        largo = any(sum(1 for c in r if _txt(c)) >= 4 for r in linhas)
        if largo:
            # cabeçalho = 1ª linha "larga" (pula um eventual título de 1 célula acima)
            h_idx = next((i for i, r in enumerate(linhas)
                          if sum(1 for c in r if _txt(c)) >= 4), 0)
            cabecalho = linhas[h_idx]
            combinacoes = []
            for row in linhas[h_idx + 1:]:
                combo = []
                for i, h in enumerate(cabecalho):
                    perg = _txt(h)
                    if not perg:
                        continue
                    val = _txt(row[i]) if i < len(row) else ""
                    if not val:
                        continue
                    combo.append({"pergunta": perg, "tipo": "", "valor": val})
                if combo:
                    combinacoes.append(combo)
            if combinacoes:
                return {"url": url, "combinacoes": combinacoes}
            # cabeçalho largo sem linhas de dados abaixo -> cai no vertical

        # VERTICAL: uma combinação (reaproveita o leitor já existente/testado)
        plano = FormsSubmitter.parse_reference_sheet(xlsx_path)
        campos = plano.get("campos", [])
        return {"url": url or plano.get("url"),
                "combinacoes": [campos] if campos else []}

    # ------------------------------------------------------------------ #
    # Valores fixos do sistema (Empresa/Coordenador/Horário sempre iguais)
    # ------------------------------------------------------------------ #
    @staticmethod
    def valores_fixos(config_path=None):
        """Lista efetiva de campos fixos: config.json (se tiver) senão os padrões.

        Nunca quebra: qualquer erro ao ler/validar o config cai nos padrões.
        """
        try:
            if config_path is None:
                from core.paths import runtime_path
                config_path = runtime_path("config.json")
            if config_path and os.path.isfile(config_path):
                with open(config_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                raw = cfg.get("forms_valores_fixos")
                campos = FormsSubmitter._coerce_fixos(raw)
                if campos:
                    return campos
        except Exception:
            pass
        return [dict(c) for c in _VALORES_FIXOS_PADRAO]

    @staticmethod
    def _coerce_fixos(raw):
        """Aceita lista de {pergunta,valor[,tipo]} OU dict {pergunta: valor}."""
        campos = []
        if isinstance(raw, dict):
            for perg, val in raw.items():
                if str(perg).strip() and str(val).strip():
                    campos.append({"pergunta": str(perg).strip(), "tipo": "",
                                   "valor": str(val).strip()})
        elif isinstance(raw, list):
            for item in raw:
                if not isinstance(item, dict):
                    continue
                perg = str(item.get("pergunta", "")).strip()
                val = str(item.get("valor", "")).strip()
                tipo = str(item.get("tipo", "")).strip()
                if perg and val:
                    campos.append({"pergunta": perg, "tipo": tipo, "valor": val})
        return campos

    @staticmethod
    def aplicar_valores_fixos(combos, fixos=None):
        """Garante os valores fixos em CADA combinação (o fixo vence o da planilha).

        Colisão por substring normalizada: um campo da planilha cujo rótulo
        contenha (ou esteja contido em) o de um fixo é removido e substituído
        pelo fixo. Assim 'Empresa', 'Nome da empresa' etc. recebem sempre o
        valor fixo, sem duplicar perguntas. Resto da planilha segue genérico.
        """
        if fixos is None:
            fixos = FormsSubmitter.valores_fixos()
        if not fixos:
            return [list(c) for c in (combos or [])]
        fixos_norm = [_norm_pergunta(f["pergunta"]) for f in fixos]

        saida = []
        for campos in (combos or []):
            base = [c for c in campos if not _colide(c.get("pergunta"), fixos_norm)]
            saida.append(base + [dict(f) for f in fixos])
        return saida

    # ------------------------------------------------------------------ #
    # Data(s) do serviço: 1 dia, intervalo (8 a 20) ou lista (7 e 8)
    # ------------------------------------------------------------------ #
    @staticmethod
    def parse_datas(texto, hoje=None):
        """Interpreta a(s) data(s) do serviço a partir do texto do usuário.

        DIA(s) — default no mês/ano correntes:
          • 1 dia:        '8'                         -> [dia 8]
          • intervalo:    '8 a 20' / '8 até 20' /
                          '8-20' / 'do dia 8 ao 20'   -> [8, 9, ..., 20]
          • lista:        '7 e 8' / '7, 8' / '8;10'   -> exatamente esses dias
          • combinação:   '8 a 10 e 15'               -> [8, 9, 10, 15]
        MÊS e ANO — opcionais (default = mês/ano de hoje):
          • por extenso:  '8 a 20 de novembro'        -> mês 11
          • por número:   '8/11' / '8 a 20/11'        -> mês 11
          • com ano:      '8/11/2026' / '8 de nov de 2027'
        Devolve lista de datetime.date ORDENADA e sem repetição; dias fora do
        mês são descartados. Texto vazio/sem dia -> []. Nunca lança.
        """
        import calendar
        from datetime import date
        if not texto or not str(texto).strip():
            return []
        hoje = hoje or date.today()
        ano, mes = hoje.year, hoje.month

        t = _norm_pergunta(texto)  # minúsculo e sem acento ('ate','ao','novembro'...)

        # ANO explícito (4 dígitos) -> captura e REMOVE p/ não ser lido como dia
        m_ano = re.search(r"\b(?:19|20)\d{2}\b", t)
        if m_ano:
            ano = int(m_ano.group(0))
            t = t[:m_ano.start()] + " " + t[m_ano.end():]

        # MÊS por extenso (pt-BR, com abreviações) -> captura e REMOVE
        _MESES = {"janeiro":1,"fevereiro":2,"marco":3,"abril":4,"maio":5,"junho":6,
                  "julho":7,"agosto":8,"setembro":9,"outubro":10,"novembro":11,
                  "dezembro":12,"jan":1,"fev":2,"mar":3,"abr":4,"mai":5,"jun":6,
                  "jul":7,"ago":8,"set":9,"out":10,"nov":11,"dez":12}
        achou_mes_nome = False
        for nome, n in sorted(_MESES.items(), key=lambda kv: -len(kv[0])):
            if re.search(r"\b" + nome + r"\b", t):
                mes = n; achou_mes_nome = True
                t = re.sub(r"\b" + nome + r"\b", " ", t)
                break
        # MÊS por número na forma com barra 'DD/MM'; sempre tira os '/NN' do texto
        barras = re.findall(r"/\s*(\d{1,2})", t)
        if barras and not achou_mes_nome and 1 <= int(barras[0]) <= 12:
            mes = int(barras[0])
        t = re.sub(r"/\s*\d{1,2}", " ", t)

        max_dia = calendar.monthrange(ano, mes)[1]
        # separa a LISTA: vírgula, ponto-e-vírgula ou a palavra 'e'
        tokens = re.split(r"\s*(?:,|;|\be\b)\s*", t)
        dias = set()
        for tok in tokens:
            nums = [int(n) for n in re.findall(r"\d+", tok)]
            if not nums:
                continue
            # intervalo? indicador de faixa ENTRE os dois primeiros números:
            # 'a'/'ao'/'ate' (palavra) ou hífen — admite palavras no meio
            # (ex.: 'do dia 8 ao dia 20'). Lista (vírgula/'e') já foi separada acima.
            faixa = len(nums) >= 2 and bool(
                re.search(r"\d+\D*?(?:\b(?:a|ao|ate)\b|[-–—])\D*?\d+", tok))
            if faixa:
                ini, fim = sorted(nums[:2])
                for d in range(ini, fim + 1):
                    dias.add(d)
            else:
                for d in nums:
                    dias.add(d)
        validos = sorted(d for d in dias if 1 <= d <= max_dia)
        return [date(ano, mes, d) for d in validos]

    @staticmethod
    def label_campo_data(config_path=None):
        """Rótulo da pergunta de DATA no formulário (default 'Data').

        Sobrescrevível por 'forms_campo_data' no config.json, sem recompilar.
        """
        try:
            if config_path is None:
                from core.paths import runtime_path
                config_path = runtime_path("config.json")
            if config_path and os.path.isfile(config_path):
                with open(config_path, "r", encoding="utf-8") as f:
                    lbl = json.load(f).get("forms_campo_data")
                if isinstance(lbl, str) and lbl.strip():
                    return lbl.strip()
        except Exception:
            pass
        return "Data"

    @staticmethod
    def _expandir_por_data(combos, datas, label=None, fmt="%d/%m/%Y"):
        """Multiplica as combinações por data: 1 preenchimento por (data × combinação).

        O campo de data vence qualquer coluna de data vinda da planilha (sem
        duplicar). Sem datas -> devolve as combinações como estão.
        """
        if not datas:
            return [list(c) for c in (combos or [])]
        label = label or FormsSubmitter.label_campo_data()
        ln = [_norm_pergunta(label)]
        valores = [d.strftime(fmt) if hasattr(d, "strftime") else str(d).strip()
                   for d in datas]
        out = []
        for v in valores:
            for campos in (combos or [[]]):
                base = [c for c in campos if not _colide(c.get("pergunta"), ln)]
                out.append(base + [{"pergunta": label, "tipo": "data", "valor": v}])
        return out

    # ------------------------------------------------------------------ #
    # Preenche o formulário a partir da planilha e gera o PDF
    # ------------------------------------------------------------------ #
    def fill_from_spreadsheet(self, xlsx_path, url=None, submit=False, log=None, timeout_ms=15000):
        """Abre o formulário no Edge, preenche/marca conforme a planilha e salva um PDF.

        Retorna {'ok', 'pdf', 'msg', 'preenchidos', 'faltas'}.
        """
        log = log or (lambda m: None)
        ok, motivo = self.is_available()
        if not ok:
            return {"ok": False, "pdf": None, "msg": f"🔒 {motivo}", "preenchidos": 0, "faltas": []}

        plano = self.parse_reference_sheet(xlsx_path)
        alvo = (url or plano.get("url") or "").strip()
        # aplica os valores fixos do sistema (Empresa/Coordenador/Horário sempre)
        campos = self.aplicar_valores_fixos([plano.get("campos", [])])[0]
        if not alvo:
            return {"ok": False, "pdf": None, "msg": "❌ Sem URL do formulário (defina 'URL' na planilha ou passe o link).", "preenchidos": 0, "faltas": []}
        if not campos:
            return {"ok": False, "pdf": None, "msg": "❌ A planilha de referência não tem campos para preencher.", "preenchidos": 0, "faltas": []}

        from playwright.sync_api import sync_playwright
        
        km_ini = next((str(c.get("valor")).strip() for c in campos if "km inicial" in str(c.get("pergunta")).lower()), None)
        km_fin = next((str(c.get("valor")).strip() for c in campos if "km final" in str(c.get("pergunta")).lower()), None)
        if not km_ini:
            km_ini = next((str(c.get("valor")).strip() for c in campos if "km" == str(c.get("pergunta")).lower() or "quilometragem" in str(c.get("pergunta")).lower()), None)

        if km_ini and km_fin:
            km_str = f"{km_ini} a {km_fin}"
        else:
            km_str = km_ini or km_fin
            
        km_safe = re.sub(r'[\\/*?:"<>|]', "", km_str) if km_str else None
        
        if km_safe:
            pdf_path = os.path.join(self.output_dir, f"KM {km_safe}.pdf")
        else:
            pdf_path = os.path.join(self.output_dir, f"Forms_{datetime.now():%Y%m%d_%H%M%S}.pdf")
        preenchidos, faltas = 0, []
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(channel="msedge", headless=True)
                page = browser.new_context().new_page()
                page.set_default_timeout(timeout_ms)
                page.emulate_media(media="print")  # = o que a opção "Imprimir" do Forms mostra
                page.goto(alvo, wait_until="domcontentloaded")
                page.wait_for_timeout(1500)
                for campo in campos:
                    try:
                        if self._preencher_campo(page, campo):
                            preenchidos += 1
                            log(f"✅ '{campo['pergunta']}' → {campo['valor'] or '(marcado)'}")
                        else:
                            faltas.append(campo["pergunta"])
                            log(f"⚠️ campo não localizado: '{campo['pergunta']}'")
                    except Exception as e:
                        faltas.append(campo["pergunta"])
                        log(f"⚠️ erro em '{campo['pergunta']}': {e}")
                page.wait_for_timeout(400)
                # "Imprimir -> Salvar como PDF": é exatamente o que page.pdf() faz
                # (renderiza a página com a mídia de impressão num PDF em disco).
                log("🖨️ Imprimindo em PDF (Imprimir → Salvar como PDF)…")
                page.pdf(path=pdf_path, print_background=True, prefer_css_page_size=True)
                if submit:
                    try:
                        page.get_by_role("button", name=re.compile(r"enviar|submit", re.I)).first.click()
                        page.get_by_text(re.compile(r"resposta foi enviada|response was submitted", re.I)).wait_for()
                        log("📨 Formulário enviado.")
                    except Exception as e:
                        log(f"⚠️ PDF gerado, mas não consegui enviar: {e}")
                browser.close()
            return {"ok": True, "pdf": pdf_path,
                    "msg": f"📄 PDF gerado ({preenchidos}/{len(campos)} campos preenchidos). Faltas: {faltas or 'nenhuma'}",
                    "preenchidos": preenchidos, "faltas": faltas}
        except Exception as e:
            return {"ok": False, "pdf": None, "msg": f"❌ Falha ao operar o Edge: {e}", "preenchidos": preenchidos, "faltas": faltas}

    # ------------------------------------------------------------------ #
    # Preenche o form UMA VEZ por combinação (vários dias) -> um PDF por combinação
    # ------------------------------------------------------------------ #
    def fill_all_from_spreadsheet(self, xlsx_path, url=None, submit=False, log=None, timeout_ms=15000, datas=None):
        """Preenche o formulário uma vez por combinação da planilha e salva um PDF
        de cada. Reaproveita o mesmo Edge para todas. O caso vertical (1 combinação)
        também passa por aqui (N=1).

        Se ``datas`` for informado (lista de date/strings OU texto cru p/ parse_datas),
        o conjunto de combinações é multiplicado por cada dia -> (data × combinação).

        Retorna {'ok','pdfs':[...],'resultados':[{comb,pdf,preenchidos,total,faltas}],'msg'}.
        """
        log = log or (lambda m: None)
        ok, motivo = self.is_available()
        if not ok:
            return {"ok": False, "pdfs": [], "resultados": [], "msg": f"🔒 {motivo}"}

        plano = self.parse_reference_sheet_multi(xlsx_path)
        alvo = (url or plano.get("url") or "").strip()
        # aplica os valores fixos do sistema em TODA combinação (sempre iguais)
        combos = self.aplicar_valores_fixos(plano.get("combinacoes", []))
        # data(s) do serviço: 1 dia, intervalo ou lista -> multiplica as combinações
        if isinstance(datas, str):
            datas = self.parse_datas(datas)
        label_data = self.label_campo_data()
        combos = self._expandir_por_data(combos, datas, label=label_data)
        if not alvo:
            return {"ok": False, "pdfs": [], "resultados": [], "msg": "❌ Sem URL do formulário (defina 'URL' na planilha ou passe o link)."}
        if not combos:
            return {"ok": False, "pdfs": [], "resultados": [], "msg": "❌ A planilha não tem combinações para preencher."}

        from playwright.sync_api import sync_playwright
        resultados = []
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(channel="msedge", headless=True)
                page = browser.new_context().new_page()
                page.set_default_timeout(timeout_ms)
                page.emulate_media(media="print")  # = o que a opção "Imprimir" do Forms mostra
                for idx, campos in enumerate(combos, 1):
                    # sufixo com a data (se houver campo de data) p/ nomear o PDF
                    ln_data = _norm_pergunta(label_data)
                    dstr = next((c.get("valor") for c in campos
                                 if _norm_pergunta(c.get("pergunta")) == ln_data), None)
                    suf = ("_" + str(dstr).replace("/", "-")) if dstr else ""
                    
                    km_ini = next((str(c.get("valor")).strip() for c in campos if "km inicial" in str(c.get("pergunta")).lower()), None)
                    km_fin = next((str(c.get("valor")).strip() for c in campos if "km final" in str(c.get("pergunta")).lower()), None)
                    if not km_ini:
                        km_ini = next((str(c.get("valor")).strip() for c in campos if "km" == str(c.get("pergunta")).lower() or "quilometragem" in str(c.get("pergunta")).lower()), None)
            
                    if km_ini and km_fin:
                        km_str = f"{km_ini} a {km_fin}"
                    else:
                        km_str = km_ini or km_fin
                        
                    km_safe = re.sub(r'[\\/*?:"<>|]', "", km_str) if km_str else None
                    
                    # Se tem data, cria subpasta só com o dia (ex: "27")
                    if dstr:
                        day_only = str(dstr).split("/")[0].strip()
                        day_folder = os.path.join(self.output_dir, day_only)
                        os.makedirs(day_folder, exist_ok=True)
                        dest_dir = day_folder
                    else:
                        dest_dir = self.output_dir
                    
                    if km_safe:
                        pdf_path = os.path.join(dest_dir, f"KM {km_safe}.pdf")
                    else:
                        pdf_path = os.path.join(dest_dir, f"Forms_{datetime.now():%Y%m%d_%H%M%S}_comb{idx:02d}.pdf")
                    preenchidos, faltas = 0, []
                    page.goto(alvo, wait_until="domcontentloaded")
                    page.wait_for_timeout(1200)
                    log(f"— combinação {idx}/{len(combos)} —")
                    for campo in campos:
                        try:
                            if self._preencher_campo(page, campo):
                                preenchidos += 1
                                log(f"✅ '{campo['pergunta']}' → {campo['valor'] or '(marcado)'}")
                            else:
                                faltas.append(campo["pergunta"])
                                log(f"⚠️ campo não localizado: '{campo['pergunta']}'")
                        except Exception as e:
                            faltas.append(campo["pergunta"])
                            log(f"⚠️ erro em '{campo['pergunta']}': {e}")
                    page.wait_for_timeout(300)
                    # "Imprimir -> Salvar como PDF" (page.pdf = mídia de impressão em disco)
                    log(f"🖨️ Imprimindo combinação {idx} em PDF…")
                    page.pdf(path=pdf_path, print_background=True, prefer_css_page_size=True)
                    if submit:
                        try:
                            page.get_by_role("button", name=re.compile(r"enviar|submit", re.I)).first.click()
                            page.get_by_text(re.compile(r"resposta foi enviada|response was submitted", re.I)).wait_for()
                            log(f"📨 combinação {idx} enviada.")
                        except Exception as e:
                            log(f"⚠️ PDF {idx} gerado, mas não consegui enviar: {e}")
                    resultados.append({"comb": idx, "pdf": pdf_path, "preenchidos": preenchidos,
                                       "total": len(campos), "faltas": faltas})
                browser.close()
            total_ok = sum(r["preenchidos"] for r in resultados)
            total_campos = sum(r["total"] for r in resultados)
            return {"ok": True, "pdfs": [r["pdf"] for r in resultados], "resultados": resultados,
                    "msg": f"📄 {len(resultados)} PDF(s) gerado(s) — {total_ok}/{total_campos} campos preenchidos no total."}
        except Exception as e:
            return {"ok": False, "pdfs": [r["pdf"] for r in resultados], "resultados": resultados,
                    "msg": f"❌ Falha ao operar o Edge: {e}"}

    # ------------------------------------------------------------------ #
    # Localiza e preenche/marca UM campo (resiliente: MS Forms e HTML genérico)
    # ------------------------------------------------------------------ #
    def _preencher_campo(self, page, campo):
        t = (campo.get("pergunta") or "").strip()
        tp = (campo.get("tipo") or "").lower().strip()
        v = (campo.get("valor") or "").strip()
        rx_t = re.compile(re.escape(t), re.I)

        # Bloco da pergunta (MS Forms novo: [data-automation-id=questionItem];
        # clássico: título -> elemento-pai). Achando o bloco, ESCOPAMOS as buscas a
        # ele — senão opções repetidas ("N/A", "SIM", "NÃO") de perguntas diferentes
        # colidiriam e marcariam a errada. Sem bloco, busca no documento (HTML genérico).
        def _question_box():
            import difflib
            best_block = None
            best_score = 0
            try:
                blocks = page.locator("[data-automation-id='questionItem'], .office-form-question").all()
                for b in blocks:
                    try:
                        title_loc = b.locator('.text-format-content, .office-form-question-title, [data-automation-id="questionTitle"]')
                        if title_loc.count():
                            title = title_loc.first.inner_text().strip()
                            clean_title = re.sub(r'^\d+[\.\)]\s*', '', title).replace('*', '').strip()
                            score = difflib.SequenceMatcher(None, t.lower(), clean_title.lower()).ratio()
                            if t.lower() in clean_title.lower():
                                score += 0.5
                            if score > best_score:
                                best_score = score
                                best_block = b
                    except Exception:
                        pass
                if best_block and best_score > 0.3:
                    return best_block
            except Exception:
                pass
            return None

        box = _question_box()
        if box is not None:
            try:
                box.scroll_into_view_if_needed(timeout=2000)  # força render de bloco virtualizado
            except Exception:
                pass

        def try_text():
            if not v:
                return False
            # 0) bloco escopado (MS Forms novo): texto, textarea ou date-combobox digitável
            if box is not None:
                try:
                    inp = box.locator("input[data-automation-id='textInput'], textarea, "
                                      "input[type='text'], input[role='combobox'], input:not([type])")
                    if inp.count():
                        inp.first.fill(v)
                        inp.first.press('Tab')
                        return True
                except Exception:
                    pass
            # 1) rótulo acessível (label for=, aria-label, aria-labelledby)
            try:
                loc = page.get_by_label(rx_t)
                if loc.count() and loc.first.is_visible():
                    el = loc.first
                    if (el.evaluate("e => e.tagName") or "").lower() == "select":
                        el.select_option(label=v)
                    else:
                        el.fill(v)
                        el.press('Tab')
                    return True
            except Exception:
                pass
            # 2) aria-label parcial
            try:
                loc = page.locator(f'input[aria-label*="{t}"], textarea[aria-label*="{t}"]')
                if loc.count():
                    loc.first.fill(v)
                    loc.first.press('Tab')
                    return True
            except Exception:
                pass
            # 3) MS Forms clássico: título da pergunta -> input dentro do bloco-pai
            try:
                b3 = page.locator('.office-form-question-title', has_text=rx_t).locator('..')
                inp = b3.locator('input[type="text"]:not([role="combobox"]), textarea, input:not([type])')
                if inp.count():
                    inp.first.fill(v)
                    inp.first.press('Tab')
                    return True
            except Exception:
                pass
            return False

        def try_choice(role):
            if not v:
                return False
            opcoes = [x.strip() for x in v.split(";") if x.strip()] or [v]
            marcou = False
            for opt in opcoes:
                feito = False
                rx_opt_exact = re.compile(f"^{re.escape(opt)}$", re.I)
                rx_opt_partial = re.compile(re.escape(opt), re.I)
                # dentro do bloco (evita colidir rótulos repetidos entre perguntas)
                if box is not None:
                    try:
                        loc = box.get_by_role(role, name=rx_opt_exact)
                        if loc.count():
                            loc.first.click(force=True)
                            feito = True
                    except Exception:
                        pass
                    if not feito:
                        try:
                            # fallback pra substring parcial
                            loc = box.get_by_role(role, name=rx_opt_partial)
                            if loc.count():
                                loc.first.click(force=True)
                                feito = True
                        except Exception:
                            pass
                    if not feito:
                        try:
                            lbl = box.locator("label").filter(has_text=rx_opt_exact)
                            if lbl.count():
                                lbl.first.click()
                                feito = True
                        except Exception:
                            pass
                    if not feito:
                        try:
                            import difflib
                            labels = box.locator("label").all_inner_texts()
                            if labels:
                                best = max(labels, key=lambda x: difflib.SequenceMatcher(None, opt.lower(), x.lower()).ratio())
                                if difflib.SequenceMatcher(None, opt.lower(), best.lower()).ratio() > 0.4:
                                    idx = labels.index(best)
                                    box.locator("label").nth(idx).click(force=True)
                                    feito = True
                        except Exception:
                            pass
                # fallback global (HTML genérico, ou bloco não localizado)
                if not feito:
                    try:
                        loc = page.get_by_role(role, name=rx_opt_exact)
                        if loc.count():
                            loc.first.click(force=True)
                            feito = True
                    except Exception:
                        pass
                if not feito:
                    try:
                        lbl = page.locator("label").filter(has_text=rx_opt_exact)
                        if lbl.count():
                            lbl.first.click()
                            feito = True
                    except Exception:
                        pass
                marcou = marcou or feito
            return marcou

        def try_dropdown():
            if not v:
                return False
            try:  # <select> nativo por rótulo
                loc = page.get_by_label(rx_t)
                if loc.count() and (loc.first.evaluate("e => e.tagName") or "").lower() == "select":
                    loc.first.select_option(label=v)
                    return True
            except Exception:
                pass
            try:  # combobox estilo MS Forms
                b2 = box or page.locator('.office-form-question-title', has_text=rx_t).locator('..')
                b2.locator('[role="combobox"], [role="listbox"], [role="button"], i[data-icon-name="ChevronDown"]').first.click()
                page.wait_for_timeout(300)
                inp = b2.locator('input[role="combobox"]')
                if inp.count():
                    inp.first.fill(v)
                    page.wait_for_timeout(300)
                page.locator('[role="option"], .ms-Dropdown-item').filter(
                    has_text=re.compile(re.escape(v), re.I)).first.click()
                return True
            except Exception:
                pass
            return False

        if tp in ("texto", "text", "data", "date", "numero", "número", "number"):
            return try_text()
        if tp in ("opcao", "opção", "radio", "escolha", "unica", "única"):
            return try_choice("radio") or try_choice("checkbox")
        if tp in ("checkbox", "caixa", "multipla", "múltipla", "marcar"):
            return try_choice("checkbox") or try_choice("radio")
        if tp in ("dropdown", "lista", "select", "combo", "combobox"):
            return try_dropdown()
        # autodetecção: texto -> radio -> checkbox -> dropdown
        return try_text() or try_choice("radio") or try_choice("checkbox") or try_dropdown()

